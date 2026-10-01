"""טבלת מצב הניקוד של כל אזורי ההתרעה (לבדיקה ולמילוי תיקונים)"""
import json, re, sys
sys.path.insert(0, 'tools'); from nikud_tools import strip
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
out = sys.argv[1]
names = json.load(open('tools/voice_names.json', encoding='utf-8'))
fix = json.load(open('tools/pronounce.json', encoding='utf-8'))
rows = []
for n, path in names['areas'].items():
    spoken = fix.get(n, n.replace(' - ', ' ').replace('-', ' '))
    words = [w for w in re.split(r'[\s,]+', spoken) if re.search('[א-ת]', w)]
    plain = [w for w in words if not any('֑' <= c <= 'ׇ' for c in w) and len(strip(w)) > 1]
    st = 'תקין' if not plain else ('חסר ניקוד' if len(plain) == len(words) else 'חלקי')
    rows.append((st, n, 'voice/' + path, spoken, ' '.join(plain)))
order = {'חסר ניקוד': 0, 'חלקי': 1, 'תקין': 2}
rows.sort(key=lambda r: (order[r[0]], r[1]))
F = 'Arial'; body = Font(name=F); bold = Font(name=F, bold=True); hdr = Font(name=F, bold=True, color='FFFFFF')
fills = {'תקין': PatternFill('solid', fgColor='D9F2DD'), 'חלקי': PatternFill('solid', fgColor='FFF2CC'), 'חסר ניקוד': PatternFill('solid', fgColor='F8D7DA')}
dark = PatternFill('solid', fgColor='333333'); yellow = PatternFill('solid', fgColor='FFFFCC')
wb = Workbook()
L = wb.active; L.title = 'הסבר'; L.sheet_view.rightToLeft = True
info = [('רשימת אזורי ההתרעה – מצב הניקוד בהקראה',), (),
        ('סטטוס', 'משמעות'),
        ('תקין', 'לכל המילים יש ניקוד – ההקראה אמורה להיות נכונה'),
        ('חלקי', 'לחלק מהמילים אין ניקוד (רשומות בעמודה "מילים בלי ניקוד")'),
        ('חסר ניקוד', 'אין ניקוד בכלל – הקול מנחש את ההגייה'), (),
        ('איך למלא', 'בגיליון "יישובים": לכתוב בעמודה הצהובה "ניקוד מתוקן" את השם המלא עם ניקוד – רק איפה שצריך תיקון – ולשלוח לי את הקובץ'),
        ('דוגמה', 'ניר עם  →  נִיר עָם'), (),
        ('סיכום',),
        ('תקין', '=COUNTIF(יישובים!D:D,"תקין")'),
        ('חלקי', '=COUNTIF(יישובים!D:D,"חלקי")'),
        ('חסר ניקוד', '=COUNTIF(יישובים!D:D,"חסר ניקוד")'),
        ('סה"כ', '=COUNTA(יישובים!B:B)-1')]
for r in info: L.append(list(r))
for row in L.iter_rows():
    for c in row: c.font = body
L['A1'].font = Font(name=F, bold=True, size=14)
for c in ('A3', 'B3', 'A8', 'A9', 'A11'): L[c].font = bold
for i, k in [(4, 'תקין'), (5, 'חלקי'), (6, 'חסר ניקוד'), (12, 'תקין'), (13, 'חלקי'), (14, 'חסר ניקוד')]: L[f'A{i}'].fill = fills[k]
L['B9'].fill = PatternFill('solid', fgColor='FFFF00')
L.column_dimensions['A'].width = 16; L.column_dimensions['B'].width = 100
W = wb.create_sheet('יישובים'); W.sheet_view.rightToLeft = True
W.append(["מס'", 'שם היישוב', 'קובץ הקול', 'סטטוס', 'מה מוקרא עכשיו', 'מילים בלי ניקוד', 'ניקוד מתוקן (למילוי)'])
for c in W[1]: c.font = hdr; c.fill = dark; c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
W['G1'].fill = PatternFill('solid', fgColor='B8860B')
line = Border(bottom=Side(style='thin', color='DDDDDD'))
for i, (st, n, f, sp, pl) in enumerate(rows, 1):
    W.append([i, n, f, st, sp, pl or None, None]); r = i + 1
    for col in 'ABCDEFG': W[f'{col}{r}'].font = body; W[f'{col}{r}'].border = line
    W[f'D{r}'].fill = fills[st]; W[f'G{r}'].fill = yellow
    W[f'C{r}'].alignment = Alignment(horizontal='left')
for col, wd in zip('ABCDEFG', [6, 30, 40, 12, 36, 26, 32]): W.column_dimensions[col].width = wd
W.freeze_panes = 'C2'; W.auto_filter.ref = f'A1:G{len(rows) + 1}'
wb.save(out)
from collections import Counter; print(Counter(r[0] for r in rows))
