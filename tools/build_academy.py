"""ניקוד לפי רשימת האקדמיה ללשון העברית (tools/academy_2022.xlsx) - גובר על הניקוד הקודם"""
import json, re, sys
sys.path.insert(0, 'tools'); from nikud_tools import strip
from openpyxl import load_workbook
ws = load_workbook('tools/academy_2022.xlsx', read_only=True).worksheets[0]
def N(s):
    s = strip(s or ''); s = re.sub(r'\(.*?\)', '', s)
    return re.sub(r'[\s\-־–,\'"״׳`]', '', s)
ac = {}
for r in list(ws.iter_rows(values_only=True))[7:]:
    if r[0] and r[1]:
        voc = re.sub(r'\s*\(.*?\)\s*', '', r[1]).strip().replace('"', '').replace('״', '')
        ac.setdefault(N(r[0]), voc)
names = json.load(open('tools/voice_names.json', encoding='utf-8'))['areas']
cur = json.load(open('tools/pronounce.json', encoding='utf-8'))
out = {}; src = {}
for n in names:
    if N(n) in ac: out[n] = ac[N(n)]; src[n] = 'academy'; continue
    parts = re.split(r'(\s+-\s*|\s*-\s+|,\s*)', n)
    if not any(N(p) in ac for p in parts[::2]):
        out[n] = cur.get(n, n); src[n] = 'claude'; continue
    # יישור מילים: המקור מול הניקוד הקודם
    ow = [w for w in re.split(r'[\s,\-]+', n) if w]
    cw = [w for w in re.split(r'[\s,]+', cur.get(n, n)) if w]
    wm = dict(zip(ow, cw)) if len(ow) == len(cw) else {}
    res = []
    for i, p in enumerate(parts):
        if i % 2: res.append(', ' if ',' in p else ' '); continue
        if N(p) in ac: res.append(ac[N(p)])
        else: res.append(' '.join(wm.get(w, w) for w in p.split()))
    out[n] = re.sub(r'\s+', ' ', ''.join(res)).strip(); src[n] = 'academy+claude'
json.dump(out, open('tools/pronounce.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
json.dump(src, open('tools/nikud_source.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
from collections import Counter; print(Counter(src.values()))
